from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches

from ..models.report import FinalReport, ReportTable


class ReportGenerationError(Exception):
    """Raised when DOCX report generation fails."""


def generate_docx(
    report: FinalReport,
    output_path: str | Path,
    include_checklist: bool = True,
) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    try:
        document = Document()

        document.add_heading(f"Stock Analysis Report: {report.company_name or report.ticker or 'Company'}", level=0)
        document.add_paragraph(f"Ticker: {report.ticker or 'N/A'} | Status: {report.status}")
        document.add_paragraph(f"Generated at: {report.generated_at.isoformat(timespec='seconds')}")

        document.add_heading("Executive Summary", level=1)

        for item in report.executive_summary:
            document.add_paragraph(item, style="List Bullet")

        for section in report.sections:
            document.add_heading(section.title, level=1)
            document.add_paragraph(section.summary)

            for bullet in section.bullets:
                document.add_paragraph(bullet, style="List Bullet")

            for table in section.tables:
                document.add_heading(table.title, level=2)
                _add_table(document, table)

        if report.chart_paths:
            document.add_heading("Technical Charts", level=1)

            for name, path in report.chart_paths.items():
                image_path = Path(path)

                if not image_path.exists():
                    continue

                document.add_heading(name.replace("_", " ").title(), level=2)
                document.add_picture(str(image_path), width=Inches(6.0))

        if report.assumptions:
            document.add_heading("Assumptions & Data Notes", level=1)

            for assumption in report.assumptions:
                document.add_paragraph(assumption, style="List Bullet")

        if include_checklist and report.checklist:
            document.add_heading("227-Point Checklist Appendix", level=1)

            checklist_table = document.add_table(rows=1, cols=5)

            try:
                checklist_table.style = "Table Grid"
            except Exception:
                pass

            header_cells = checklist_table.rows[0].cells
            header_cells[0].text = "Item"
            header_cells[1].text = "Section"
            header_cells[2].text = "Description"
            header_cells[3].text = "Module"
            header_cells[4].text = "Status"

            for item in report.checklist:
                row_cells = checklist_table.add_row().cells
                row_cells[0].text = item.item_id
                row_cells[1].text = item.section
                row_cells[2].text = item.description
                row_cells[3].text = item.module
                row_cells[4].text = item.status

        document.save(str(output))
        return output
    except Exception as exc:
        raise ReportGenerationError(f"Failed to generate DOCX report: {output}") from exc


def _add_table(document: Document, table: ReportTable) -> None:
    doc_table = document.add_table(rows=1, cols=len(table.headers))

    try:
        doc_table.style = "Table Grid"
    except Exception:
        pass

    header_cells = doc_table.rows[0].cells

    for index, header in enumerate(table.headers):
        header_cells[index].text = header

    for row in table.rows:
        row_cells = doc_table.add_row().cells

        for index, cell in enumerate(row):
            row_cells[index].text = "" if cell is None else str(cell)
